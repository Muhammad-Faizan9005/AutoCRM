import sys
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from fastapi.testclient import TestClient

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from app.auth.utils import create_access_token
from app.database import get_db
from app.main import app
from app.middleware.rate_limiter import reset_rate_limiter_state


class FakeResponse:
    def __init__(self, data: Any):
        self.data = data


class FakeTable:
    def __init__(self, table_name: str, db: "FakeDB"):
        self.table_name = table_name
        self.db = db
        self._filters: list[tuple[str, Any]] = []
        self._single = False
        self._insert_payload: dict[str, Any] | None = None
        self._update_payload: dict[str, Any] | None = None
        self._is_delete = False

    def select(self, *_args, **_kwargs):
        return self

    def eq(self, column: str, value: Any):
        self._filters.append((column, value))
        return self

    def limit(self, _value: int):
        return self

    def single(self):
        self._single = True
        return self

    def range(self, *_args, **_kwargs):
        return self

    def order(self, *_args, **_kwargs):
        return self

    def insert(self, payload: dict[str, Any]):
        self._insert_payload = payload
        return self

    def update(self, payload: dict[str, Any]):
        self._update_payload = payload
        return self

    def delete(self):
        self._is_delete = True
        return self

    def execute(self):
        rows = self.db.tables[self.table_name]

        def matches(row: dict[str, Any]) -> bool:
            return all(str(row.get(col)) == str(val) for col, val in self._filters)

        if self._insert_payload is not None:
            row = self._insert_payload.copy()
            row.setdefault("id", str(uuid4()))
            row.setdefault("created_at", datetime.now(timezone.utc).isoformat())
            row.setdefault("updated_at", datetime.now(timezone.utc).isoformat())
            self.db.tables[self.table_name].append(row)
            return FakeResponse([row.copy()])

        if self._update_payload is not None:
            updated_rows: list[dict[str, Any]] = []
            for row in rows:
                if matches(row):
                    row.update(self._update_payload)
                    row["updated_at"] = datetime.now(timezone.utc).isoformat()
                    updated_rows.append(row.copy())
            return FakeResponse(updated_rows)

        if self._is_delete:
            deleted_rows = [row.copy() for row in rows if matches(row)]
            self.db.tables[self.table_name] = [row for row in rows if not matches(row)]
            return FakeResponse(deleted_rows)

        filtered = [row.copy() for row in rows if matches(row)]
        if self._single:
            return FakeResponse(filtered[0] if filtered else None)
        return FakeResponse(filtered)


class FakeDB:
    def __init__(self):
        now = datetime.now(timezone.utc).isoformat()

        self.admin_id = str(uuid4())
        self.manager_id = str(uuid4())
        self.rep_id = str(uuid4())
        self.existing_customer_id = str(uuid4())

        self.tables: dict[str, list[dict[str, Any]]] = {
            "agents": [
                {
                    "id": self.admin_id,
                    "email": "admin@example.com",
                    "full_name": "Admin User",
                    "role": "admin",
                    "is_active": True,
                    "password_hash": "hashed::x",
                    "created_at": now,
                    "updated_at": now,
                },
                {
                    "id": self.manager_id,
                    "email": "manager@example.com",
                    "full_name": "Manager User",
                    "role": "sales_manager",
                    "is_active": True,
                    "password_hash": "hashed::x",
                    "created_at": now,
                    "updated_at": now,
                },
                {
                    "id": self.rep_id,
                    "email": "rep@example.com",
                    "full_name": "Rep User",
                    "role": "sales_rep",
                    "is_active": True,
                    "password_hash": "hashed::x",
                    "created_at": now,
                    "updated_at": now,
                },
            ],
            "customers": [
                {
                    "id": self.existing_customer_id,
                    "email": "existing@example.com",
                    "full_name": "Existing Customer",
                    "phone": "+1 555 1000",
                    "company": "Old Co",
                    "status": "active",
                    "notes": "Old note",
                    "created_at": now,
                    "updated_at": now,
                }
            ],
            "tickets": [],
            "ticket_messages": [],
            "revoked_tokens": [],
        }

    def table(self, table_name: str) -> FakeTable:
        if table_name not in self.tables:
            self.tables[table_name] = []
        return FakeTable(table_name, self)


def _token_for(user_id: str) -> str:
    return create_access_token({"sub": user_id})


def _client_with_fake_db() -> tuple[TestClient, FakeDB]:
    fake_db = FakeDB()
    app.dependency_overrides[get_db] = lambda: fake_db
    reset_rate_limiter_state()
    return TestClient(app), fake_db


def _build_excel_bytes(rows: list[dict[str, str]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active

    headers = list(rows[0].keys())
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_import_customers_from_csv_creates_and_updates():
    client, fake_db = _client_with_fake_db()
    manager_headers = {"Authorization": f"Bearer {_token_for(fake_db.manager_id)}"}

    csv_content = "\n".join(
        [
            "email,full_name,phone,company,status,notes",
            "existing@example.com,<b>Existing Updated</b>,+1 555 2000,New Co,lead,Updated note",
            "new@example.com,New Customer,+1 555 3000,Acme Inc,active,Fresh lead",
        ]
    )

    response = client.post(
        "/api/import/customers",
        headers=manager_headers,
        files={"file": ("customers.csv", csv_content.encode("utf-8"), "text/csv")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["entity"] == "customers"
    assert payload["total_rows"] == 2
    assert payload["created_count"] == 1
    assert payload["updated_count"] == 1
    assert payload["failed_count"] == 0

    updated = next(row for row in fake_db.tables["customers"] if row["email"] == "existing@example.com")
    assert updated["full_name"] == "Existing Updated"
    assert updated["status"] == "lead"
    assert len(fake_db.tables["customers"]) == 2


def test_import_tickets_from_excel_with_row_level_failure():
    client, fake_db = _client_with_fake_db()
    manager_headers = {"Authorization": f"Bearer {_token_for(fake_db.manager_id)}"}

    excel_bytes = _build_excel_bytes(
        [
            {
                "customer_email": "existing@example.com",
                "subject": "Need onboarding help",
                "description": "Cannot complete onboarding",
                "status": "open",
                "priority": "high",
                "category": "support",
            },
            {
                "customer_email": "missing@example.com",
                "subject": "Missing customer",
                "description": "Should fail",
                "status": "open",
                "priority": "low",
                "category": "support",
            },
        ]
    )

    response = client.post(
        "/api/import/tickets",
        headers=manager_headers,
        files={"file": ("tickets.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["entity"] == "tickets"
    assert payload["total_rows"] == 2
    assert payload["created_count"] == 1
    assert payload["failed_count"] == 1
    assert payload["failures"][0]["row_number"] == 3
    assert len(fake_db.tables["tickets"]) == 1


def test_import_requires_manager_or_admin_role():
    client, fake_db = _client_with_fake_db()
    rep_headers = {"Authorization": f"Bearer {_token_for(fake_db.rep_id)}"}

    csv_content = "email,full_name\nrepnew@example.com,Rep Should Fail"
    response = client.post(
        "/api/import/customers",
        headers=rep_headers,
        files={"file": ("customers.csv", csv_content.encode("utf-8"), "text/csv")},
    )

    assert response.status_code == 403


def test_import_rejects_unsupported_file_type():
    client, fake_db = _client_with_fake_db()
    manager_headers = {"Authorization": f"Bearer {_token_for(fake_db.manager_id)}"}

    response = client.post(
        "/api/import/customers",
        headers=manager_headers,
        files={"file": ("customers.txt", b"invalid", "text/plain")},
    )

    assert response.status_code == 400
    assert "Unsupported file type" in response.json()["error"]["message"]
